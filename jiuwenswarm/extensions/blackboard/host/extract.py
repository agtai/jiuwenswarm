"""The text of a reference file, for agents (milestone 7): plain text, Markdown, CSV and the like as
they are, PDF through pdfplumber, Word through python-docx. Done on the host, so every member's agent
reads the same text.
"""

from __future__ import annotations

from pathlib import Path

from jiuwenswarm.extensions.blackboard.common.errors import BlackboardError

MAX_CHARS = 200_000
UNSUPPORTED = "unsupported_reference"

TEXT_SUFFIXES = {".md", ".markdown", ".txt", ".csv", ".tsv", ".json", ".yaml", ".yml", ".xml", ".log", ".html", ".htm"}
TEXT_MIMES = ("text/", "application/json", "application/xml", "application/x-yaml", "application/yaml")
PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def kind_of(name: str, mime: str) -> str:
    suffix = Path(name).suffix.lower()
    if mime.startswith(TEXT_MIMES) or suffix in TEXT_SUFFIXES:
        return "text"
    if mime == PDF_MIME or suffix == ".pdf":
        return "pdf"
    if mime == DOCX_MIME or suffix == ".docx":
        return "docx"
    if mime.startswith("image/"):
        return "image"
    return "other"


def _pdf(path: Path) -> str:
    import pdfplumber

    parts: list[str] = []
    size = 0
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            parts.append(text)
            size += len(text)
            if size > MAX_CHARS:
                break
    return "\n\n".join(parts)


def _docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        lines.append("")
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def extract_text(path: Path, name: str, mime: str) -> dict:
    """``{kind, text, truncated}``; images and other files raise ``unsupported_reference``."""
    kind = kind_of(name, mime)
    try:
        if kind == "text":
            with path.open("rb") as f:
                text = f.read(MAX_CHARS * 4 + 1).decode("utf-8", errors="replace")
        elif kind == "pdf":
            text = _pdf(path)
        elif kind == "docx":
            text = _docx(path)
        elif kind == "image":
            raise BlackboardError(UNSUPPORTED, "images cannot be read as text", {"kind": "image", "mime": mime})
        else:
            raise BlackboardError(UNSUPPORTED, f"{mime or 'this'} files cannot be read as text", {"kind": kind, "mime": mime})
    except BlackboardError:
        raise
    except Exception as exc:  # noqa: BLE001 - a damaged file is the reader's problem, not the host's
        raise BlackboardError(UNSUPPORTED, f"the file could not be read: {exc}", {"kind": kind}) from exc
    return {"kind": kind, "text": text[:MAX_CHARS], "truncated": len(text) > MAX_CHARS}
